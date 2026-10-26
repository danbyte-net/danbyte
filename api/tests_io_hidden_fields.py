"""Spreadsheet export shows a user the fields the type's API shows them, and
no more: a field a serializer hides per user (the SLA service credit without
``view_credits``) is neither a column nor a usable export filter (#387)."""
from __future__ import annotations

import csv
import io
import json
from decimal import Decimal

from django.contrib.auth.models import User
from rest_framework.test import APITestCase

from auth_api.models import ObjectPermission, UserProfile
from core.models import Organization, Tenant

CREDIT = ("credit_tiers", "period_fee", "currency")


class _Case(APITestCase):
    def setUp(self):
        from monitoring.models import SlaAgreement, SlaTemplate

        org = Organization.objects.create(name="O", slug="o")
        self.tenant = Tenant.objects.create(org=org, name="T", slug="t")
        credits = {
            "period_fee": Decimal("12345.67"), "currency": "EUR",
            "credit_tiers": [{"below": 99.5, "credit_pct": 25}],
        }
        self.template = SlaTemplate.objects.create(
            tenant=self.tenant, name="Gold", target_pct="99.9", **credits)
        self.agreement = SlaAgreement.objects.create(
            tenant=self.tenant, name="Cust", target_pct=Decimal("99"), **credits)

    def _user(self, name, grants):
        u = User.objects.create_user(name, password="x")
        UserProfile.objects.create(user=u, role="custom").tenants.add(self.tenant)
        for slug, actions in grants.items():
            perm = ObjectPermission.objects.create(
                name=f"{name}-{slug}", object_types=[slug], actions=list(actions),
            )
            perm.users.add(u)
            perm.tenants.add(self.tenant)
        return u

    def _login(self, user):
        self.client.force_login(user)
        s = self.client.session
        s["current_tenant_id"] = str(self.tenant.id)
        s.save()

    def _csv(self, slug, query=""):
        resp = self.client.get(f"/api/io/{slug}/export/?fmt=csv{query}")
        self.assertEqual(resp.status_code, 200, getattr(resp, "content", b""))
        text = b"".join(resp.streaming_content).decode()
        reader = csv.DictReader(io.StringIO(text))
        return reader.fieldnames, list(reader)


class HiddenCreditTests(_Case):
    def setUp(self):
        super().setUp()
        self.viewer = self._user("viewer", {
            "slatemplate": ["view"], "slaagreement": ["view"],
        })
        self.finance = self._user("finance", {
            "slatemplate": ["view"], "slaagreement": ["view", "view_credits"],
        })

    def test_api_hides_credits(self):
        """The baseline the export must match."""
        self._login(self.viewer)
        body = self.client.get(f"/api/monitoring/sla-templates/{self.template.id}/").json()
        self.assertFalse(set(CREDIT) & set(body))

    def test_csv_export_leaves_credits_out(self):
        self._login(self.viewer)
        for slug in ("slatemplate", "slaagreement"):
            cols, rows = self._csv(slug)
            self.assertEqual(len(rows), 1, slug)
            self.assertFalse(set(CREDIT) & set(cols), slug)
            self.assertNotIn("12345.67", json.dumps(rows))

    def test_json_and_xlsx_exports_leave_credits_out(self):
        from openpyxl import load_workbook

        self._login(self.viewer)
        for slug in ("slatemplate", "slaagreement"):
            resp = self.client.get(f"/api/io/{slug}/export/?fmt=json")
            rows = json.loads(b"".join(resp.streaming_content))
            self.assertFalse(set(CREDIT) & set(rows[0]), slug)
            resp = self.client.get(f"/api/io/{slug}/export/?fmt=xlsx")
            ws = load_workbook(io.BytesIO(resp.content)).active
            header = [c.value for c in next(ws.iter_rows(max_row=1))]
            self.assertFalse(set(CREDIT) & set(header), slug)
            self.assertNotIn("12345.67", json.dumps(
                [[c.value for c in r] for r in ws.iter_rows()], default=str))

    def test_post_export_leaves_credits_out(self):
        self._login(self.viewer)
        resp = self.client.post("/api/io/slatemplate/export/",
                                {"fmt": "json", "ids": [str(self.template.id)]},
                                format="json")
        rows = json.loads(b"".join(resp.streaming_content))
        self.assertFalse(set(CREDIT) & set(rows[0]))

    def test_hidden_field_is_refused_as_a_filter(self):
        self._login(self.viewer)
        for slug in ("slatemplate", "slaagreement"):
            for f, val in (("period_fee", "12345.67"), ("currency", "EUR")):
                resp = self.client.get(f"/api/io/{slug}/export/?fmt=csv&{f}={val}")
                self.assertEqual(resp.status_code, 400, (slug, f))
            resp = self.client.post(f"/api/io/{slug}/export/",
                                    {"fmt": "csv", "period_fee": "12345.67"},
                                    format="json")
            self.assertEqual(resp.status_code, 400, slug)

    def test_fields_endpoint_leaves_credits_out(self):
        self._login(self.viewer)
        body = self.client.get("/api/io/slatemplate/fields/").json()
        self.assertFalse(set(CREDIT) & set(body["columns"]))
        self.assertFalse(set(CREDIT) & {f["name"] for f in body["fields"]})

    def test_view_credits_still_exports_and_filters_them(self):
        self._login(self.finance)
        for slug in ("slatemplate", "slaagreement"):
            cols, rows = self._csv(slug)
            self.assertTrue(set(CREDIT) <= set(cols), slug)
            self.assertEqual(Decimal(rows[0]["period_fee"]), Decimal("12345.67"))
            self.assertEqual(len(self._csv(slug, "&period_fee=12345.67")[1]), 1)
            self.assertEqual(len(self._csv(slug, "&period_fee=1.00")[1]), 0)

    def test_visible_fields_still_filter(self):
        self._login(self.viewer)
        self.assertEqual(len(self._csv("slatemplate", "&name=Gold")[1]), 1)
        self.assertEqual(len(self._csv("slatemplate", "&name=Other")[1]), 0)

    def test_round_trip_without_credits_keeps_them(self):
        """A file exported without the credit columns imports back without
        touching the stored credits."""
        editor = self._user("editor", {"slatemplate": ["view", "change"]})
        self._login(editor)
        _cols, rows = self._csv("slatemplate")
        rows[0]["description"] = "edited"
        resp = self.client.post("/api/io/slatemplate/import/",
                                {"rows": rows, "dry_run": False}, format="json")
        self.assertEqual(resp.json()["errors"], [])
        self.template.refresh_from_db()
        self.assertEqual(self.template.description, "edited")
        self.assertEqual(self.template.period_fee, Decimal("12345.67"))
        self.assertEqual(self.template.currency, "EUR")

    def test_secret_store_paths_are_refused_as_filters(self):
        """A column export never writes is never a filter either."""
        self.superuser = User.objects.create_user("root", password="x", is_superuser=True)
        UserProfile.objects.create(user=self.superuser, role="admin").tenants.add(self.tenant)
        self._login(self.superuser)
        resp = self.client.get("/api/io/wirelesslan/export/?fmt=csv&psk_secret_path=x")
        self.assertEqual(resp.status_code, 400)


class SerializerHidingIsDeclaredTests(APITestCase):
    """Every serializer behind an export type that changes its output per
    request either declares the fields it hides (``hidden_fields``) or is
    known not to hide any - so a new per-user field drop cannot slip past
    the export unnoticed."""

    #: to_representation overrides that reshape values, never drop by user.
    RESHAPE_ONLY = {
        "HolidayCalendarSerializer",   # normalises legacy day strings
        "OuiPrefixSerializer",         # prefix in display form
        "SlaIncidentFollowUpSerializer",  # flat dict, same fields for all
        # Drops the channel secrets for everyone alike and answers *_set
        # flags instead; nothing is hidden per user.
        "NotificationChannelSerializer",
    }

    def test_every_override_is_accounted_for(self):
        from rest_framework import serializers

        from api.io import io_types, read_viewset_for
        from auth_api.object_types import model_for
        from core.hidden_fields import PermissionHiddenFieldsMixin

        unexplained = []
        for t in io_types():
            vs = read_viewset_for(model_for(t["slug"]))
            ser_cls = getattr(vs, "serializer_class", None) if vs else None
            if ser_cls is None:
                continue
            overrides = any(
                "to_representation" in vars(k) for k in ser_cls.__mro__
                if k not in (serializers.ModelSerializer, serializers.Serializer,
                             serializers.BaseSerializer, serializers.Field)
                and not k.__module__.startswith("rest_framework")
            )
            if (overrides and not issubclass(ser_cls, PermissionHiddenFieldsMixin)
                    and ser_cls.__name__ not in self.RESHAPE_ONLY):
                unexplained.append(ser_cls.__name__)
        self.assertEqual(unexplained, [])

    def test_sla_serializers_declare_their_credit_fields(self):
        from core.hidden_fields import PermissionHiddenFieldsMixin
        from monitoring.sla_api import SlaAgreementSerializer
        from monitoring.sla_template_api import SlaTemplateSerializer

        for cls in (SlaAgreementSerializer, SlaTemplateSerializer):
            self.assertTrue(issubclass(cls, PermissionHiddenFieldsMixin), cls)
